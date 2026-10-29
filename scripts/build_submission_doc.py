# -*- coding: utf-8 -*-
"""生成作品介绍文档（预指导提交用）"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SHOTS = r"D:\Projects\_shots"
OUT = r"D:\Projects\ai-study-assistant\智页_AI智能学习助手_作品介绍文档.docx"

doc = Document()

# ---------- 全局样式 ----------
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(6)

for i in range(1, 4):
    hs = doc.styles[f"Heading {i}"]
    hs.font.name = "微软雅黑"
    hs.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    hs.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

# 页面设置
sec = doc.sections[0]
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

def h(text, level=1):
    p = doc.add_heading(text, level=level)
    return p

def para(text, bold=False, size=10.5, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    for r in p.runs:
        r.font.name = "微软雅黑"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        r.font.size = Pt(10.5)
    return p

def numbered(idx, text):
    """手动编号，避免 Word 列表跨章节连续编号"""
    p = doc.add_paragraph()
    r1 = p.add_run(f"{idx}. ")
    r1.font.name = "微软雅黑"
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r1.font.size = Pt(10.5)
    r2 = p.add_run(text)
    r2.font.name = "微软雅黑"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r2.font.size = Pt(10.5)
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.9)
    p.paragraph_format.space_after = Pt(6)
    return p

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, txt in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(txt)
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(9.5)
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, txt in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            run = p.add_run(str(txt))
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            run.font.size = Pt(9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t

def image(name, caption, width_cm=14.8):
    path = os.path.join(SHOTS, name)
    if not os.path.exists(path):
        para(f"[缺少图片：{name}]", size=9)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))
    # 图片与其标题保持在同一页，避免标题孤零零留在下一页页首
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(6)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    cap.paragraph_format.space_after = Pt(14)

# ==================== 封面 ====================
for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("AI智能学习助手")
run.font.name = "微软雅黑"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
run.font.size = Pt(34)
run.bold = True
run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("作 品 介 绍 文 档")
run.font.name = "微软雅黑"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

for _ in range(4):
    doc.add_paragraph()

info = [
    ("参赛赛事", "第九届传智杯全国大学生数智创新与AI应用大赛"),
    ("参赛赛道", "AI WEB 网页开发"),
    ("参赛组别", "本科 B 组"),
    ("团队名称", "智页"),
    ("所在学校", "岭南师范学院"),
    ("作品名称", "AI智能学习助手"),
    ("访问地址", "https://ai-study-assistant-2aig.onrender.com"),
]
t = doc.add_table(rows=0, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for k, v in info:
    cells = t.add_row().cells
    cells[0].text = ""
    r1 = cells[0].paragraphs[0].add_run(k)
    r1.font.name = "微软雅黑"
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r1.font.size = Pt(11)
    r1.bold = True
    cells[0].width = Cm(3.2)
    cells[1].text = ""
    r2 = cells[1].paragraphs[0].add_run(v)
    r2.font.name = "微软雅黑"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r2.font.size = Pt(11)
    cells[1].width = Cm(11.3)
    for c in cells:
        for pp in c.paragraphs:
            pp.paragraph_format.space_before = Pt(2)
            pp.paragraph_format.space_after = Pt(2)
            pp.paragraph_format.line_spacing = 1.15

doc.add_page_break()

# ==================== 一、项目详情及使用方式 ====================
h("一、项目详情及使用方式", 1)

h("1.1 项目简介", 2)
para(
    "AI智能学习助手是一款面向在校学生的个性化学习辅助平台，围绕“学、练、测、复习、诊断”的完整学习闭环设计，"
    "提供AI答疑、知识点拆解、学习计划、错题本、艾宾浩斯复习、学习看板、AI诊断七大核心功能。"
)
para(
    "与传统AI问答产品的根本区别在于：本系统实现了检索增强生成（RAG）技术。当学生提问时，系统会自动检索该学生本人的错题本、"
    "知识点拆解、学习计划与AI诊断报告，将检索结果作为上下文注入大模型，使AI的回答建立在该学生的真实学情之上，"
    "从而给出真正个性化的辅导，而不是通用化的泛泛而谈。"
)
para(
    "同时，系统集成多模态视觉理解能力，支持拍照录题：学生拍摄错题照片上传，AI自动识别科目、题目内容、错误答案与正确答案，"
    "一次性回填表单，大幅降低错题录入的机械成本。"
)

h("1.2 功能概览", 2)
table(
    ["功能模块", "核心能力", "技术实现"],
    [
        ["AI智能答疑", "多轮对话解答各学科问题，MathJax级公式渲染", "大模型 + RAG检索增强 + 自研LaTeX渲染器"],
        ["知识点拆解", "从定义、要点、举例、误区、延伸五维结构化拆解", "大模型结构化生成 + 知识图谱可视化"],
        ["学习计划", "根据科目、考试日期、每日时长生成按天规划", "大模型规划生成 + 导出功能"],
        ["错题本", "拍照录题、AI错因分析、类似题生成、三级掌握度", "多模态视觉识别 + 大模型分析"],
        ["艾宾浩斯复习", "基于遗忘曲线自动安排复习间隔", "间隔重复算法 + 复习状态机"],
        ["学习数据看板", "热力图、雷达图、趋势图、饼图四类可视化", "自研SVG图表引擎"],
        ["AI学习诊断", "综合分析学情生成结构化诊断报告", "大模型 + 多维数据聚合"],
    ],
    widths=[3.2, 6.5, 5.8]
)

h("1.3 使用方式", 2)
para("线上访问", bold=True)
para("访问地址：https://ai-study-assistant-2aig.onrender.com")
para(
    "说明：线上部署于免费云服务，连续15分钟无访问会进入休眠，首次打开需等待约30至60秒冷启动，属正常现象。"
)
para("测试账号", bold=True)
table(
    ["用户名", "密码", "说明"],
    [["111", "123456", "张三，含完整示例数据（错题、知识点、计划、诊断报告）"],
     ["testuser", "123456", "基础测试账号"]],
    widths=[3.5, 3.5, 8.5]
)
para("各功能使用步骤", bold=True)
numbered(1, "注册登录：首页点击“免费开始”，填写用户名、密码与安全问题（用于找回密码），注册后自动登录。")
numbered(2, "AI答疑：在输入框提问，系统自动检索个人学习数据后作答，回答上方出现“已基于你的学习数据回答”标识即表示RAG生效。")
numbered(3, "知识点拆解：输入知识点名称（可指定科目），点击“开始拆解”，可在“列表”与“知识图谱”两种视图间切换。")
numbered(4, "学习计划：填写科目、考试日期、每日学习时长，点击生成，支持导出为文本文件。")
numbered(5, "拍照录题：进入错题本 → 添加错题 → 上传错题照片 → 点击“AI识别题目”，科目、题目、答案自动填充。")
numbered(6, "艾宾浩斯复习：错题本顶部横幅显示今日待复习数量，点击“开始复习”逐题复习并标记掌握程度。")
numbered(7, "学习看板：查看热力图、雷达图、趋势图与错题科目分布饼图。")
numbered(8, "AI诊断：点击“生成我的诊断报告”，综合分析学情后输出结构化报告，支持打印或导出PDF。")

h("1.4 本地运行方式", 2)
para("本项目已完成云端部署，可直接访问线上地址体验。如需在本地运行，步骤如下：")
numbered(1, "环境要求：Node.js 18.0 及以上版本。")
numbered(2, "执行 npm install 安装依赖。")
numbered(3, "复制 .env.example 为 .env，填入大模型API密钥与推理接入点ID。")
numbered(4, "执行 npm start 启动服务。")
numbered(5, "浏览器访问 http://localhost:3000。")
para(
    "说明：项目默认使用本地 SQLite 数据库，零配置即可运行；生产环境可通过环境变量切换为 PostgreSQL 与云端对象存储。"
)

doc.add_page_break()

# ==================== 二、项目亮点 ====================
h("二、项目亮点", 1)

h("2.1 核心亮点：RAG个性化智能答疑", 2)
para(
    "这是本作品与普通AI问答产品最本质的区别。普通产品中，同一道问题问100个人，得到的是同一种回答；"
    "而在本系统中，AI会先“翻看”提问者自己的学习档案，再作答。"
)
para("技术链路如下：", bold=True)
para(
    "用户提问 → 问题向量化 → 与个人学习数据向量做余弦相似度计算 → 取Top相关记录 → 组装为上下文注入提示词 → "
    "大模型生成个性化回答 → 返回并标识“已基于你的学习数据回答”"
)
para("实际效果", bold=True)
para(
    "当学生提问“我数学有哪些薄弱点”时，AI不会给出笼统的学科建议，而是引用其错题本中的具体题目。实测输出中，"
    "AI准确识别出“一元一次方程移项时符号处理错误”“二次函数区间最值的分类讨论不完整”等具体薄弱点，"
    "并结合诊断报告给出针对性练习建议。"
)
para("工程保障：双通道检索", bold=True)
para(
    "为保证功能在任何环境下都可用，系统设计了语义检索与关键词匹配两条通道。当向量服务不可用时（未开通、模型退役、网络异常），"
    "系统自动降级为学科关键词匹配，并缓存该状态避免重复失败。两条通道都能检索到个人数据并注入上下文，"
    "用户侧体验一致，功能不中断。"
)

h("2.2 差异化亮点：多模态拍照录题", 2)
para(
    "错题录入是学生使用学习工具时最大的摩擦点——手抄一道题平均需要1至2分钟。本系统支持直接拍照上传，"
    "AI自动识别科目、题目内容、错误答案与正确答案四类信息并回填表单，将录入成本降低到“拍一张照”的程度。"
)
para(
    "技术上使用火山方舟的多模态模型的视觉理解能力，实测单张图片5至15秒完成识别，对印刷体与手写体均有效。"
    "视觉模型与文本模型解耦配置，可独立替换为任意支持图片输入的模型。"
)

h("2.3 知识图谱可视化", 2)
para(
    "将分散的知识点、错题、学习计划组织成一张关系网络，支持环形布局与树形布局两种展示方式。"
    "节点按科目配色，错题节点用颜色深浅表示掌握程度，点击节点可直接跳转到对应的知识点或错题详情，"
    "让“知识之间有什么关联”一目了然，而不是散落在各自列表里。"
)

h("2.4 艾宾浩斯遗忘曲线复习", 2)
para(
    "错题不是记下来就结束了。系统基于遗忘曲线算法为每道错题安排复习时间，每天自动计算并提示当日待复习内容，"
    "复习后根据掌握程度动态调整下次复习间隔，形成“记录—复习—巩固”的闭环，避免一错再错。"
)

h("2.5 自研零依赖 LaTeX 公式渲染器", 2)
para(
    "AI答疑涉及大量数学公式，若依赖第三方渲染库会增加体积与加载时间。项目自研了轻量级公式渲染器，"
    "支持行内公式与块级公式，内置100多个数学符号的Unicode映射与上下标自动转换，零外部依赖，"
    "在AI答疑、知识点拆解、错题分析、诊断报告等所有模块自动生效。"
)

h("2.6 自研SVG数据可视化", 2)
para(
    "学习看板的四类图表（学习热力图、知识点掌握雷达图、学习时长趋势图、错题科目分布饼图）以及知识图谱，"
    "全部采用原生SVG手写实现，未引入任何图表库。优势是体积小、可完全控住样式、深色模式适配彻底，"
    "并且能针对学习场景做定制化设计。"
)

h("2.7 工程亮点：一套代码，两种部署形态", 2)
para(
    "项目在部署阶段遇到一个典型工程问题：免费云平台的容器文件系统是临时的，服务休眠或重启后本地数据库与上传图片会被清空。"
)
para(
    "解决方案不是重写业务代码，而是抽象出数据库适配层与存储适配层：数据库方面，业务代码继续使用统一的查询接口，"
    "由适配层自动完成 SQL 方言转换（占位符、自增主键、冲突处理等），本地用 SQLite、生产用 PostgreSQL；"
    "图片方面，存储适配层统一封装上传、读取、存在性检查与删除，本地落磁盘、生产传云端对象存储。"
)
para(
    "最终通过两个环境变量即可切换部署形态，应用服务器变成无状态节点，可随时重启、重新部署或迁移平台，数据不受影响。"
    "实测服务休眠16分钟后冷启动，数据库记录与图片文件均完整保留。"
)

doc.add_page_break()

# ==================== 三、AI技术落地使用情况 ====================
h("三、AI技术落地使用情况概述", 1)

h("3.1 模型选型与分工", 2)
para("项目接入火山方舟（Volcengine Ark）平台，采用双模型分工架构：")
table(
    ["用途", "模型", "配置项", "表现"],
    [
        ["文本任务：答疑、知识点拆解、学习计划、诊断、错题分析", "DeepSeek-V4.1-Flash", "DOUBAO_MODEL", "单次响应 3 至 11 秒"],
        ["视觉任务：拍照识别错题", "Doubao-Seed-Evolving（多模态）", "VISION_MODEL", "单张识别 5 至 15 秒"],
        ["语义检索：文本向量化", "doubao-embedding", "EMBEDDING_MODEL", "不可用时自动降级"],
    ],
    widths=[6.2, 4.3, 3.2, 3.0]
)
para(
    "选型说明：DeepSeek 系列为纯文本模型，不具备图片输入能力，因此视觉任务必须使用独立的多模态模型。"
    "两个模型通过环境变量独立配置，任一模型都可在不改代码的情况下替换。"
)

h("3.2 RAG检索增强生成的落地细节", 2)
para("检索范围", bold=True)
para("错题本、知识点拆解、学习计划、AI诊断报告四类个人数据。")
para("向量化策略", bold=True)
para(
    "采用懒加载：历史数据按需向量化并缓存到向量表，不一次性全量处理，兼顾响应速度与调用成本。"
)
para("检索与注入", bold=True)
para(
    "计算问题向量与各条学习数据的余弦相似度，按相似度排序取Top记录，再按来源分类（错题、知识点、计划、诊断），"
    "格式化为结构化文本注入提示词。每个类别设有条数上限，避免上下文过长。"
)
para("降级与容错", bold=True)
para(
    "向量服务不可用时自动切换关键词匹配，并在首次失败时缓存状态，后续不再重复尝试，避免无效请求与日志噪音。"
)

h("3.3 多模态识别的落地细节", 2)
para("处理流程", bold=True)
para(
    "图片上传（内存缓冲）→ 转Base64 → 调用多模态模型（结构化提示词要求返回JSON）→ 解析JSON（兼容Markdown代码块包裹）"
    "→ 校验识别结果 → 回填表单"
)
para("识别结果校验", bold=True)
para(
    "为避免“识别失败却提示成功”的问题，后端对结果做了归一化与有效性判断：仅当识别出题目内容时才标记为成功，"
    "否则明确返回失败，前端提示用户重拍或手动填写。"
)
para("提示词优化", bold=True)
para(
    "针对手写作业场景，在提示词中明确告知图片可能是手机拍摄、字迹潦草、存在阴影或倾斜，要求模型尽力辨认，"
    "并将“只要能看到题目就必须输出”作为硬性约束。"
)

h("3.4 AI能力使用场景统计", 2)
table(
    ["AI能力", "使用场景", "调用频次（典型使用）"],
    [
        ["大模型文本生成", "答疑、知识点拆解、学习计划、错题分析、诊断报告", "每次用户操作 1 次"],
        ["多模态视觉理解", "拍照录题", "每次录题 1 次"],
        ["文本向量嵌入", "RAG语义检索", "每次答疑 1 次（可降级）"],
    ],
    widths=[4.5, 8.0, 4.5]
)

h("3.5 成本与稳定性控制", 2)
bullet("双模型分工：文本任务用成本较低的对话模型，仅视觉任务调用多模态模型。")
bullet("懒加载向量化：避免全量重复计算。")
bullet("上下文裁剪：RAG注入的每类数据设有条数上限，控制提示词长度。")
bullet("降级兜底：任一外部服务异常都不阻塞主流程，功能可用性优先。")

doc.add_page_break()

# ==================== 四、关键技术 ====================
h("四、项目涉及的关键技术", 1)

h("4.1 技术栈总览", 2)
table(
    ["层次", "技术选型"],
    [
        ["前端", "原生 HTML5 / CSS3 / JavaScript（零框架、零CDN、全部资源本地化）"],
        ["后端", "Node.js + Express 4.x"],
        ["数据库", "SQLite（本地开发） / PostgreSQL（生产部署），自研适配层自动转换方言"],
        ["图片存储", "本地磁盘（开发） / 云端对象存储（生产），自研存储适配层"],
        ["AI能力", "火山方舟 大语言模型 + 多模态模型 + 文本向量模型"],
        ["认证", "JWT + bcryptjs 密码加密"],
        ["文件上传", "multer（内存缓冲模式）"],
    ],
    widths=[3.0, 13.0]
)

h("4.2 系统架构", 2)
para("系统分为四层：", bold=True)
para(
    "表现层（浏览器）→ 接口层（Express 路由）→ 能力层（AI封装、RAG、存储适配、数据库适配）→ 数据层（数据库与对象存储）"
)
para("接口层模块划分：", bold=True)
table(
    ["模块", "职责"],
    [
        ["routes/auth.js", "注册、登录、忘记密码、用户信息"],
        ["routes/chat.js", "AI对话、会话与消息管理，集成RAG"],
        ["routes/knowledge.js", "知识点拆解、知识图谱数据"],
        ["routes/plan.js", "学习计划生成与导出"],
        ["routes/mistake.js", "错题增删改查、拍照识别、艾宾浩斯复习"],
        ["routes/dashboard.js", "看板统计：概览、热力图、雷达图、趋势图、科目分布"],
        ["routes/diagnosis.js", "AI诊断报告生成与管理"],
        ["routes/quiz.js", "AI出题与答题记录"],
        ["routes/upload.js", "图片上传与存储维护"],
    ],
    widths=[4.5, 11.5]
)
para("能力层模块划分：", bold=True)
table(
    ["模块", "职责"],
    [
        ["utils/doubao.js", "大模型调用封装，含思考型模型参数控制与重试兜底"],
        ["utils/vision.js", "多模态识别封装，含结果校验"],
        ["utils/embedding.js", "向量嵌入封装，含服务可用性缓存"],
        ["utils/rag.js", "RAG检索增强生成，语义检索与关键词降级双通道"],
        ["utils/storage.js", "图片存储适配层，本地磁盘 / 云端对象存储"],
        ["db/adapter.js", "数据库适配层，SQL方言自动转换"],
    ],
    widths=[4.5, 11.5]
)

h("4.3 数据库设计", 2)
para("共10张表：")
table(
    ["表名", "用途"],
    [
        ["users", "用户信息（含安全问题，用于找回密码）"],
        ["conversations / messages", "AI对话会话与消息"],
        ["knowledge_points", "知识点拆解记录（含科目标签）"],
        ["study_plans", "学习计划"],
        ["mistakes", "错题本（含图片路径、掌握度、艾宾浩斯复习字段）"],
        ["study_records", "每日学习记录（用于热力图与趋势图）"],
        ["diagnosis_reports", "AI诊断报告"],
        ["quiz_sets", "AI出题记录"],
        ["knowledge_embeddings", "向量缓存（语义检索用）"],
    ],
    widths=[5.0, 11.0]
)

h("4.4 关键技术点说明", 2)
para("数据库适配层", bold=True)
para(
    "业务代码统一使用 db.prepare(sql).get()/.all()/.run() 接口，适配层内部完成：位置占位符转换为PostgreSQL的编号占位符、"
    "具名参数重排、自增主键语法转换、冲突处理语法转换、时间与数值类型映射，以及插入后返回自增主键的兼容处理。"
    "同时统一数值类型，避免PostgreSQL的bigint被驱动转为字符串后导致前端计算出错。"
)
para("存储适配层", bold=True)
para(
    "上传路由使用内存缓冲接收文件，由存储层决定落盘还是上传云端；图片引用在数据库中统一存文件名，"
    "前端按存储模式解析为完整地址。视觉识别模块也通过存储层读取图片内容，因此两种模式下识别功能都能正常工作。"
)
para("前端工程化实践", bold=True)
para(
    "零框架、零CDN，所有第三方资源（字体、Marked、Highlight）本地化，离线可用；资源版本号管理避免浏览器缓存；"
    "响应式适配桌面与移动端；深浅色主题切换时图表自动重绘。"
)

doc.add_page_break()

# ==================== 五、项目图片 ====================
h("五、项目相关图片", 1)
para("以下为系统实际运行截图。")

image("01-落地页.png", "图1  产品落地页（首屏与功能展示）")
image("02-AI答疑.png", "图2  AI智能答疑（基于个人学习数据的个性化回答）")
image("03-知识点拆解.png", "图3  知识点拆解（五维结构化）")
image("08-知识图谱.png", "图4  知识图谱可视化（环形布局）")
image("04-学习计划.png", "图5  学习计划生成")
image("05-错题本.png", "图6  错题本管理")
image("06-学习看板.png", "图7  学习数据看板（热力图与雷达图）")
image("07-AI诊断.png", "图8  AI学习诊断")
image("09-深色模式-看板.png", "图9  深色模式适配")

doc.add_page_break()

# ==================== 六、发表/获奖情况 ====================
h("六、项目发表或获奖情况说明", 1)
para(
    "本作品尚未在任何平台发表，也未在其他比赛中获奖。本次系首次参加第九届传智杯全国大学生数智创新与AI应用大赛。"
)
para(
    "作品已完成线上部署并对外开放访问，访问地址：https://ai-study-assistant-2aig.onrender.com，"
    "评审专家可直接在线体验全部功能。"
)
para(
    "源代码已在公开仓库托管，便于评审核查："
)
para(
    "GitHub：https://github.com/adragonsr911-ctrl/ai-study-assistant",
)
para(
    "Gitee：https://gitee.com/sr911/ai-study-assistant",
)

# ==================== 七、原创性说明 ====================
h("七、原创性与优化说明", 1)
para("本作品不是在其他项目基础上优化的产物，而是从零开始自主设计与开发的原创作品。", bold=True)

h("7.1 团队自主完成的部分", 2)
bullet("整体产品设计与功能规划，围绕学习闭环设计七大功能模块。")
bullet("全部前后端业务代码：路由层、能力层、数据层、前端交互逻辑。")
bullet("RAG检索策略：检索范围、向量化时机、上下文组织方式与降级方案。")
bullet("数据库与存储适配层：SQL方言转换规则、双模式切换机制。")
bullet("自研LaTeX公式渲染器：符号映射表与解析逻辑。")
bullet("自研SVG可视化组件：热力图、雷达图、趋势图、饼图、知识图谱的绘制算法与布局算法。")
bullet("知识图谱的环形/树形布局算法与节点跳转交互。")
bullet("艾宾浩斯复习的间隔调度逻辑。")

h("7.2 第三方资源使用说明", 2)
table(
    ["资源", "用途", "许可"],
    [
        ["Marked.js", "Markdown渲染", "MIT"],
        ["Highlight.js", "代码高亮", "BSD"],
        ["Lucide Icons", "线性图标库（本地化44个图标）", "ISC"],
        ["Inter 字体", "界面字体（本地化）", "SIL OFL"],
        ["multer", "文件上传处理", "MIT"],
        ["Express / better-sqlite3 / pg 等", "后端基础依赖", "MIT"],
    ],
    widths=[6.0, 6.5, 3.5]
)
para(
    "以上第三方资源均遵循其开源许可协议使用，且已在项目中本地化存储，不依赖外部CDN。"
)

h("7.3 AI工具使用说明", 2)
para(
    "项目在开发过程中使用了AI辅助编程工具协助代码编写与调试。核心的产品设计、技术选型、架构决策、"
    "算法实现与问题排查均由团队成员主导完成，团队成员能够清晰解释全部代码逻辑。"
)
para(
    "需要特别说明的是，项目中的AI能力（大模型生成、多模态识别、向量检索）是通过调用合规的云端API实现的，"
    "API密钥通过环境变量配置，未硬编码在公开代码仓库中。"
)

# 保存
doc.save(OUT)
print("文档已生成:", OUT)
print("大小:", round(os.path.getsize(OUT) / 1024 / 1024, 2), "MB")
